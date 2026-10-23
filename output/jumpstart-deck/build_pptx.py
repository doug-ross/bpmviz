"""Build the AIxRay Jump Start deck on the Sogeti April 2026 PowerPoint template.

Usage: python -I build_pptx.py <template.pptx> <deck-folder> <out.pptx>
The deck folder is output/jumpstart-deck (assets and the HTML deck with speaker notes).
"""
import copy, json, re, sys
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

TEMPLATE, DECK, OUT = sys.argv[1], Path(sys.argv[2]), sys.argv[3]
A = DECK / "assets"

# Sogeti April 2026 theme
NAVY = RGBColor(0x12, 0x1A, 0x38)
BLUE = RGBColor(0x00, 0x58, 0xAB)
VIBRANT = RGBColor(0x1D, 0xB8, 0xF2)
RED = RGBColor(0xF5, 0x3B, 0x53)
BODY = RGBColor(0x34, 0x34, 0x3D)
MUTED = RGBColor(0x6B, 0x6B, 0x78)
RULE = RGBColor(0xDD, 0xDD, 0xE2)
RULE_STRONG = RGBColor(0xC5, 0xC5, 0xCD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT_WHITE = RGBColor(0xD9, 0xDC, 0xE6)
HEAD, LIGHT = "Ubuntu", "Ubuntu Light"

notes_html = (DECK / "AIxRay-Jump-Start.html").read_text(encoding="utf-8")
NOTES = json.loads(re.search(r'<script type="application/json" id="speaker-notes">(.*?)</script>', notes_html, re.S).group(1))

prs = Presentation(TEMPLATE)
M0, M1 = prs.slide_masters[0], prs.slide_masters[1]
LAY_LIGHT = M0.slide_layouts.get_by_name("Only Title - Grey")
LAY_DARK = M1.slide_layouts.get_by_name("Only Title - Grey")

# ---- keep the template's title, closing-section, about and logo slides; drop the rest ----
sld_ids = prs.slides._sldIdLst
orig = list(sld_ids)
KEEP = {0: "title", 4: "close", 17: "about", 18: "logo"}
kept = {}
for i, sid in enumerate(orig):
    if i in KEEP:
        kept[KEEP[i]] = sid
    else:
        prs.part.drop_rel(sid.rId)
        sld_ids.remove(sid)
slide_by_id = {s.slide_id: s for s in prs.slides}
S_TITLE, S_CLOSE, S_ABOUT, S_LOGO = (slide_by_id[int(kept[k].get("id"))] for k in ("title", "close", "about", "logo"))


# ---------------- helpers ----------------
from PIL import ImageFont
_FONT_FILES = {(LIGHT, False): "Ubuntu-L.ttf", (LIGHT, True): "Ubuntu-B.ttf", (HEAD, False): "Ubuntu-R.ttf", (HEAD, True): "Ubuntu-B.ttf"}
_font_cache = {}


def _font(name, bold):
    key = (name, bold)
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype("C:/Windows/Fonts/" + _FONT_FILES[key], 200)
    return _font_cache[key]


def measure(paras, w, size=14, font=LIGHT, bold=False, line=1.15, space_after=0):
    """Height in inches PowerPoint needs for the paragraphs, wrapped greedily at word boundaries."""
    if isinstance(paras, str):
        paras = [paras]
    width_pt = w * 72 * 0.97
    total = 0.0
    for para in paras:
        runs = [(para, {})] if isinstance(para, str) else para
        words = []
        for text, o in runs:
            f = _font(o.get("font", font), o.get("bold", bold))
            sz = o.get("size", size)
            for wd in re.findall(r"\S+\s*", text):
                words.append((f.getlength(wd.rstrip()) * sz / 200, f.getlength(wd) * sz / 200))
        lines, cur = 1, 0.0
        for bare, full in words:
            if cur and cur + bare > width_pt:
                lines, cur = lines + 1, 0.0
            cur += full
        total += lines * size * 1.2 * line / 72 + space_after / 72
    return total
def set_text(shape, text):
    """Replace a placeholder's text, keeping the first run's formatting. '\v' becomes a line break."""
    tf = shape.text_frame
    paras = tf.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    p = paras[0]
    runs = p.runs
    rpr = copy.deepcopy(runs[0]._r.find(qn("a:rPr"))) if runs else None
    for child in list(p._p):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p._p.remove(child)
    end = p._p.find(qn("a:endParaRPr"))
    for k, part in enumerate(text.split("\v")):
        if k:
            br = etree.SubElement(p._p, qn("a:br"))
            if end is not None:
                end.addprevious(br)
        r = etree.SubElement(p._p, qn("a:r"))
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = etree.SubElement(r, qn("a:t"))
        t.text = part
        if end is not None:
            end.addprevious(r)


def tb(slide, x, y, w, h, paras, size=14, color=BODY, font=LIGHT, bold=False, anchor=MSO_ANCHOR.TOP,
       space_after=0, line=1.15, align=PP_ALIGN.LEFT, name=None):
    """Text box. paras: list of paragraphs; each a string or a list of (text, {overrides}) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for text, o in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.size = Pt(o.get("size", size))
            f.name = o.get("font", font)
            f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color)
    return box


def hline(slide, x, y, w, color=RULE, weight=0.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def card(slide, x, y, w, h, border=RULE, weight=0.75, fill=WHITE):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = border
    sh.line.width = Pt(weight)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def picture(slide, path, x, y, w=None, h=None, frame=False):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    if frame:
        pic.line.color.rgb = RULE
        pic.line.width = Pt(0.75)
    return pic


def eyebrow(slide, text, dark=False):
    box = tb(slide, 0.66, 1.18, 9, 0.3, text.upper(), size=11, color=VIBRANT if dark else BLUE, font=HEAD, bold=True)
    for r in box.text_frame.paragraphs[0].runs:
        r.font._rPr.set("spc", "120")
    return box


def new_slide(title, label, dark=False):
    s = prs.slides.add_slide(LAY_DARK if dark else LAY_LIGHT)
    s.shapes.title.text = title
    eyebrow(s, label, dark)
    return s


def rows(slide, x, y, w, items, num_w=0.55, row_h=None, size=14, num_size=18, gap=0.16, color=BODY, rule=RULE, num_color=BLUE):
    """Numbered rows separated by hairlines. items: (marker, bold lead or None, text)."""
    cy = y
    for k, (marker, lead, text) in enumerate(items):
        runs = ([(lead + " ", {"font": HEAD, "bold": True, "color": NAVY if rule != SOFT_RULE else WHITE})] if lead else []) + [(text, {})]
        h = max(measure([runs], w - num_w, size), 0.32)
        tb(slide, x, cy, num_w, 0.4, marker, size=num_size, color=num_color, font=HEAD, bold=True)
        tb(slide, x + num_w, cy + 0.02, w - num_w, h, [runs], size=size, color=color)
        cy += h + gap
        if k < len(items) - 1:
            hline(slide, x, cy - gap / 2 - 0.0, w, rule)
            cy += 0.0
    return cy


SOFT_RULE = RGBColor(0x3A, 0x42, 0x63)


def caption(slide, x, y, w, text, color=MUTED):
    return tb(slide, x, y, w, 0.3, text, size=11, color=color)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------- 01 title (template slide) ----------------
for sh in list(S_TITLE.shapes):
    if sh.is_placeholder and sh.placeholder_format.idx == 0:
        set_text(sh, "See the process. Reimagine the possibilities.")
    elif sh.is_placeholder and sh.placeholder_format.idx == 1:
        set_text(sh, "How we use AIxRay in a Jump Start workshop to understand how your work runs today and design how it could run with supervised AI.")
    elif not sh.is_placeholder and sh.left >= prs.slide_width:
        sh._element.getparent().remove(sh._element)   # spare classification badges and their instruction, parked off-slide
notes(S_TITLE, NOTES[0])

# ---------------- 02 executive summary ----------------
s = new_slide("One shared view of a process turns AI ideas into decisions", "Executive summary")
tb(s, 0.66, 1.75, 4.6, 4.6, [
    "Many AI conversations stall because nobody can point to where the work happens, who does it and what a change would be worth.",
    "In a Jump Start, we put one of your processes under AIxRay and work through it with the people who run it. The same model carries the discussion from today's process to a future design and a first estimate of its value.",
], size=17, color=NAVY, space_after=14, line=1.2)
rows(s, 5.75, 1.8, 6.92, [
    ("01", "Understand.", "We build a model of today's process from your documents and the people in the room, then play real cases through it, exceptions included."),
    ("02", "Reimagine.", "We design a future version, trying simpler fixes first and using supervised AI where it earns its place, with approvals and fallbacks written down."),
    ("03", "Evaluate.", "We estimate cycle time, effort and cost from assumptions everyone can see and challenge, and agree what to validate next."),
    ("→", "You leave with", "an agreed current-state model, a future-state design, a prioritized list of opportunities and a pilot brief with named owners."),
], row_h=0.78, gap=0.3)
notes(s, NOTES[1])

# ---------------- 03 the challenge ----------------
s = new_slide("AI ideas need a real process to land in", "Why start with the process")
tb(s, 0.66, 1.72, 10.5, 0.7, "Before deciding where AI fits, a leadership team needs the same picture of how work moves today. Three things usually get in the way.", size=17, color=NAVY)
cols = [
    ("chat.png", "Process knowledge sits with people", "How work really moves, where it waits and how exceptions are handled often lives in people's heads and in documents that disagree with each other."),
    ("cube.png", "AI opportunities stay abstract", "Without a shared view of the steps, it is hard to say where AI helps, where a simpler fix would do, and where people need to stay in control."),
    ("dollar.png", "Business cases hide their assumptions", "Benefit figures are hard to trust when nobody can see the volumes, effort and costs behind them, or which of those were measured and which were guessed."),
]
for k, (icon, head, body) in enumerate(cols):
    x = 0.66 + k * 4.12
    picture(s, A / "icons-3d" / icon, x, 2.85, h=1.0)
    hh = measure(head, 3.7, 17, HEAD, True)
    tb(s, x, 4.1, 3.7, hh, head, size=17, color=NAVY, font=HEAD, bold=True)
    tb(s, x, 4.1 + hh + 0.12, 3.7, 1.8, body, size=14)
notes(s, NOTES[2])

# ---------------- 04 what AIxRay is ----------------
s = new_slide("AIxRay is a workbench for process workshops", "The tool")
rows(s, 0.66, 1.8, 4.55, [
    ("01", None, "Plays a case through the process step by step, on one screen the whole room can follow."),
    ("02", None, "Records pain points, questions and ideas on the step where they come up."),
    ("03", None, "Compares today's process with a future design on the same clock."),
    ("04", None, "Shows every figure with its formula and inputs, labeled by how much we know about them."),
    ("05", None, "Runs in the browser from local files, with no server, account or installation."),
], num_w=0.5, row_h=[0.5, 0.5, 0.5, 0.5, 0.5], gap=0.36, num_size=16)
picture(s, A / "screens" / "02-process-step.png", 5.6, 1.8, w=7.07, frame=True)
caption(s, 5.6, 6.32, 7.07, "Illustrative example: commercial property underwriting, demonstration data.")
notes(s, NOTES[3])

# ---------------- 05 how it fits ----------------
s = new_slide("One model carries the work from understanding to evaluation", "How AIxRay fits in a Jump Start")
tb(s, 0.66, 1.72, 11, 0.4, "The workshop runs as three conversations, understand, reimagine and evaluate, with preparation before and a clear hand-over after.", size=17, color=NAVY)
stages = [
    ("BEFORE", "Prepare", "We gather your procedures, maps and notes, plus a system event log where one exists, and draft a first model.", "AI-drafted model, reviewed by our team", False),
    ("01", "Understand", "We walk routine cases and exceptions with the people who do the work, correct the model live and capture findings.", "Walkthrough and findings at each step", True),
    ("02", "Reimagine", "We try removing, simplifying and connecting steps before adding automation and AI, and design the controls.", "Future alternatives and automation designs", True),
    ("03", "Evaluate", "We model cycle time, effort and cost against today's baseline, test key assumptions and rank opportunities.", "Value, sensitivity, matrix and waves", True),
    ("AFTER", "Hand over", "We agree a checkpoint, produce the pilot brief and report, and send a review form to wider stakeholders.", "Pilot brief, report, review form and exports", False),
]
cw, gap = 2.27, 0.165
for k, (lab, head, body, screen, core) in enumerate(stages):
    x = 0.66 + k * (cw + gap)
    hline(s, x, 2.55, cw, BLUE if core else RULE_STRONG, 2.25)
    tb(s, x, 2.72, cw, 0.3, lab, size=11, color=BLUE if core else MUTED, font=HEAD, bold=True)
    tb(s, x, 3.0, cw, 0.4, head, size=17, color=NAVY, font=HEAD, bold=True)
    tb(s, x, 3.42, cw, 1.9, body, size=13.5)
    tb(s, x, 5.92, cw, 0.6, screen, size=12.5, color=NAVY, font=HEAD)
hline(s, 0.66, 5.45, 12.01, RULE)
eb = tb(s, 0.66, 5.58, 6, 0.3, "ON SCREEN IN AIXRAY", size=10.5, color=MUTED, font=HEAD, bold=True)
eb.text_frame.paragraphs[0].runs[0].font._rPr.set("spc", "120")
notes(s, NOTES[4])

# ---------------- 06 understand ----------------
s = new_slide("Agree how the work runs today", "01  Understand")
picture(s, A / "screens" / "03-capture.png", 0.66, 1.8, w=6.75, frame=True)
caption(s, 0.66, 6.1, 6.75, "Findings recorded on the step they belong to. Illustrative example, demonstration data.")
rows(s, 7.85, 1.8, 4.82, [
    ("A", "Start from what you have.", "Procedures, maps, slides, notes or BPMN files. AI can draft the first model; our team reviews it before you see it."),
    ("B", "Walk the routes.", "The routine case first, then the exceptions, choosing each branch with the participants."),
    ("C", "Capture at the step.", "Pain points, open questions and decisions are recorded where they arise, each with an owner."),
    ("D", "Keep evidence apart from belief.", "Each step shows where it came from and whether you have confirmed it."),
], num_w=0.45, row_h=[0.95, 0.72, 0.72, 0.72], gap=0.26)
notes(s, NOTES[5])

# ---------------- 07 reimagine ----------------
s = new_slide("Design the future process against a locked baseline", "02  Reimagine")
rows(s, 0.66, 1.8, 4.6, [
    ("A", "Simpler changes first.", "Each opportunity asks whether removing, simplifying, standardizing or connecting systems would do the job before automation or AI."),
    ("B", "A baseline that cannot drift.", "Today's process is locked. Every future alternative records which steps changed, were added or removed, and why."),
    ("C", "Side by side, on one clock.", "Both versions play together, so the room sees where time, handoffs and waits fall away."),
], num_w=0.45, row_h=[1.2, 1.2, 0.95], gap=0.3)
picture(s, A / "screens" / "05-compare.png", 5.75, 1.8, w=6.92, frame=True)
caption(s, 5.75, 6.2, 6.92, "Current and proposed future state with the scorecard. Illustrative example, demonstration data.")
notes(s, NOTES[6])

# ---------------- 08 controls ----------------
s = new_slide("Every AI step is designed with its controls", "02  Reimagine")
tb(s, 0.66, 1.8, 3.9, 3.2, [
    "For each step where automation or AI does the work, we write down how it is supervised before it goes anywhere near a pilot.",
    "Approvals, permissions, escalation and test criteria stay attached to the design and flow into the pilot brief and its acceptance criteria.",
], size=16, color=NAVY, space_after=12, line=1.2)
picture(s, A / "icons-3d" / "thumbs-up.png", 0.66, 5.05, h=0.85)
tb(s, 1.55, 5.2, 3.0, 0.6, "People keep the decisions that need judgment.", size=14, color=NAVY, font=HEAD, bold=True)

spec = [
    ("Design element", "Example: AI checks application completeness"),
    ("Trigger", "Application received"),
    ("Rules and limits", "Check required fields only. Do not approve applications."),
    ("Approval", "Case handler approves every external message."),
    ("Fallback and escalation", "Return to manual intake when extraction fails. Escalate uncertain cases to the duty manager."),
    ("Test criteria", "Flag every missing required document in the agreed test set."),
    ("Owner", "Operations lead"),
]
heights = [0.45, 0.5, 0.5, 0.5, 0.78, 0.5, 0.5]
gt = s.shapes.add_table(len(spec), 2, Inches(5.0), Inches(1.8), Inches(7.67), Inches(sum(heights)))
tbl = gt.table
tbl.first_row = False
tbl.horz_banding = False
tbl.columns[0].width = Inches(2.35)
tbl.columns[1].width = Inches(5.32)


def cell_lines(cell, bottom_color, bottom_w):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        tcPr.remove(child)
    for side in ("a:lnL", "a:lnR", "a:lnT"):
        ln = etree.SubElement(tcPr, qn(side), w="0")
        etree.SubElement(ln, qn("a:noFill"))
    ln = etree.SubElement(tcPr, qn("a:lnB"), w=str(int(Pt(bottom_w))), cap="flat", cmpd="sng", algn="ctr")
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=str(bottom_color))
    etree.SubElement(tcPr, qn("a:noFill"))


for r, (a, b) in enumerate(spec):
    tbl.rows[r].height = Inches(heights[r])
    for c, text in enumerate((a, b)):
        cell = tbl.cell(r, c)
        cell.margin_left = 0
        cell.margin_right = Inches(0.15)
        cell.margin_top = Inches(0.08)
        cell.margin_bottom = Inches(0.08)
        cell.vertical_anchor = MSO_ANCHOR.TOP
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = text.upper() if r == 0 else text
        f = run.font
        if r == 0:
            f.size, f.name, f.bold, f.color.rgb = Pt(10.5), HEAD, True, MUTED
            run.font._rPr.set("spc", "80")
        elif c == 0:
            f.size, f.name, f.bold, f.color.rgb = Pt(14), HEAD, True, NAVY
        else:
            f.size, f.name, f.bold, f.color.rgb = Pt(14), LIGHT, False, BODY
        cell_lines(cell, "C5C5CD" if r == 0 else "DDDDE2", 1.0 if r == 0 else 0.75)
caption(s, 5.0, 1.8 + sum(heights) + 0.15, 7.67, "Illustrative example from a synthetic workshop.")
notes(s, NOTES[7])

# ---------------- 09 evaluate ----------------
s = new_slide("Directional estimates you can check line by line", "03  Evaluate")
measures = [("Cycle time", "Start to finish, waits included"), ("Human effort", "Hands-on hours per case"),
            ("Operating cost", "Labor plus AI and service cost"), ("Freed capacity", "Time released, in FTE")]
for k, (m, d) in enumerate(measures):
    x = 0.66 + (k % 2) * 2.45
    y = 1.8 + (k // 2) * 0.95
    tb(s, x, y, 2.3, 0.35, m, size=17, color=BLUE, font=HEAD, bold=True)
    tb(s, x, y + 0.38, 2.3, 0.35, d, size=12.5)
    if k < 2:
        hline(s, x, y + 0.86, 2.25)
rows(s, 0.66, 3.85, 4.75, [
    ("→", None, "Every figure shows its formula and inputs. Each input is labeled illustrative, provided, measured or confirmed."),
    ("→", None, "A missing input leaves the result blank rather than assuming zero."),
    ("→", None, "A sensitivity view shows which inputs move the result most, which tells us what to validate first."),
], num_w=0.38, row_h=[0.72, 0.5, 0.72], gap=0.2, num_size=14)
picture(s, A / "screens" / "08-value.png", 5.85, 1.8, w=6.82, frame=True)
caption(s, 5.85, 6.1, 6.82, "Directional estimates for decision-making, not guaranteed results. Illustrative example, demonstration data.")
notes(s, NOTES[8])

# ---------------- 10 prioritize ----------------
s = new_slide("From a list of findings to a plan in waves", "03  Evaluate")
tb(s, 0.66, 1.72, 11, 0.7, "Each opportunity carries a small business case: effort saved, cases affected and cost to deliver. That lets us compare them on the same terms.", size=17, color=NAVY)
tb(s, 0.66, 2.62, 5.5, 0.4, "Value against effort", size=17, color=NAVY, font=HEAD, bold=True)
qx, qy, qw, qh, g = 1.05, 3.4, 2.55, 1.15, 0.12
tb(s, qx, 3.08, qw, 0.3, "Lower effort", size=11, color=MUTED, align=PP_ALIGN.CENTER)
tb(s, qx + qw + g, 3.08, qw, 0.3, "Higher effort", size=11, color=MUTED, align=PP_ALIGN.CENTER)
for row_i, lab in enumerate(("Higher value", "Lower value")):
    t = tb(s, 0.5, qy + row_i * (qh + g) + qh / 2 - 0.15, qh, 0.3, lab, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    t.rotation = 270
    t.left = Inches(0.66 + 0.15 - qh / 2)
quads = [("Quick wins", "Start here"), ("Major projects", "Plan and fund deliberately"),
         ("Fill-ins", "Do when capacity allows"), ("Reconsider", "Challenge before committing")]
for k, (h_, d_) in enumerate(quads):
    x = qx + (k % 2) * (qw + g)
    y = qy + (k // 2) * (qh + g)
    card(s, x, y, qw, qh, BLUE if k == 0 else RULE, 2.0 if k == 0 else 0.75)
    tb(s, x + 0.22, y + 0.22, qw - 0.4, 0.35, h_, size=15, color=NAVY, font=HEAD, bold=True)
    tb(s, x + 0.22, y + 0.6, qw - 0.4, 0.4, d_, size=12.5)
tb(s, 7.0, 2.62, 5.5, 0.4, "Delivery waves", size=17, color=NAVY, font=HEAD, bold=True)
rows(s, 7.0, 3.15, 5.67, [
    ("Now", None, "Quick wins and the first supervised AI pilot, with acceptance criteria and owners."),
    ("Next", None, "Changes that depend on the pilot's results or on connecting systems."),
    ("Later", None, "Larger redesigns, kept visible with their open questions."),
], num_w=0.9, row_h=[0.5, 0.5, 0.5], gap=0.3, num_size=16)
tb(s, 7.0, 5.65, 5.67, 0.6, "Each wave shows a running total of its modeled value, built from the same formulas as the individual opportunities.", size=12.5, color=MUTED)
notes(s, NOTES[9])

# ---------------- 11 measured data ----------------
s = new_slide("Add measured data when you have it", "Beyond the workshop")
tb(s, 0.66, 1.72, 11, 0.7, "Workshop estimates are a starting point. If your workflow, ticketing or core system can export an event log, the same model can show how the process actually runs.", size=17, color=NAVY)
for k, (icon, h_, d_) in enumerate([
    ("growth.png", "Measure today's process", "How many cases reach each step, how long each step takes and which paths cases follow, shown next to what the workshop captured."),
    ("cycle.png", "Check whether a change worked", "A log from before a change and one from after are compared step by step, with caveats such as seasonality and volume mix stated alongside."),
]):
    x = 0.66 + k * 4.25
    card(s, x, 2.75, 3.95, 3.3)
    picture(s, A / "icons-3d" / icon, x + 0.3, 3.0, h=0.75)
    hh = measure(h_, 3.4, 17, HEAD, True)
    tb(s, x + 0.3, 3.95, 3.4, hh, h_, size=17, color=NAVY, font=HEAD, bold=True)
    tb(s, x + 0.3, 3.95 + hh + 0.12, 3.4, 1.5, d_, size=14)
tb(s, 9.3, 2.85, 3.37, 0.4, "Only aggregates are kept", size=15, color=NAVY, font=HEAD, bold=True)
tb(s, 9.3, 3.27, 3.37, 1.4, "Counts and times per step. Individual events, case IDs and the names of people in the log are not stored.", size=14)
notes(s, NOTES[10])

# ---------------- 12 deliverables ----------------
s = new_slide("What you take away from the Jump Start", "Outcomes")
items = [
    ("Current-state model", "Agreed with the people who run the process, with its sources and open questions."),
    ("Future-state design", "One or more alternatives compared with the baseline, each change with its reason."),
    ("Prioritized opportunities", "Ranked by value against effort and grouped into now, next and later."),
    ("Pilot brief", "Automation designs with their controls, acceptance criteria and owned actions."),
    ("Report and review form", "A printable summary, and a form that lets stakeholders confirm or dispute each step."),
    ("Exports for your tools", "A PowerPoint deck, a BPMN 2.0 diagram, and a Jira, Azure DevOps or Excel backlog."),
]
cw, ch, gx, gy = 3.85, 2.05, 0.23, 0.3
for k, (h_, d_) in enumerate(items):
    x = 0.66 + (k % 3) * (cw + gx)
    y = 1.85 + (k // 3) * (ch + gy)
    card(s, x, y, cw, ch)
    tb(s, x + 0.3, y + 0.25, 1, 0.3, f"{k + 1:02d}", size=12, color=BLUE, font=HEAD, bold=True)
    hh = measure(h_, cw - 0.6, 17, HEAD, True)
    tb(s, x + 0.3, y + 0.58, cw - 0.6, hh, h_, size=17, color=NAVY, font=HEAD, bold=True)
    tb(s, x + 0.3, y + 0.58 + hh + 0.1, cw - 0.6, 0.9, d_, size=14)
notes(s, NOTES[11])

# ---------------- 13 data handling (dark) ----------------
s = new_slide("Your material stays under your control", "Data handling", dark=True)
blocks = [
    ("Runs from local files", "Opened from disk, AIxRay makes no network requests. The saved project file is the record, stored like any other client document."),
    ("AI drafting is optional and labeled", "AI runs only on approved commercial plans with the project lead's approval. Every AI draft stays out of reports and exports until a person reviews it."),
    ("Sensitive findings stay out by default", "Findings marked sensitive are left out of reports, review forms and exports unless someone deliberately includes them."),
    ("Review forms carry only what you choose", "Steps and selected findings, with no figures, sources or notes. The form sends nothing anywhere; answers come back as a file."),
]
for k, (h_, d_) in enumerate(blocks):
    x = 0.66 + (k % 2) * 6.1
    y = 1.85 + (k // 2) * 1.75
    tb(s, x, y, 5.6, 0.4, h_, size=17, color=WHITE, font=HEAD, bold=True)
    tb(s, x, y + 0.45, 5.6, 1.1, d_, size=14, color=SOFT_WHITE)
hline(s, 0.66, 5.45, 11.8, SOFT_RULE)
tb(s, 0.66, 5.62, 11.5, 0.6, "Where material includes personal or otherwise regulated data, we agree the handling with your data protection officer before the workshop.", size=14, color=SOFT_WHITE)
notes(s, NOTES[12])

# ---------------- 14 what we need ----------------
s = new_slide("What we ask of you before the workshop", "Getting ready")
rows(s, 0.66, 1.85, 7.3, [
    ("01", "A candidate process", "with a clear start, a clear end and a business owner who wants it improved."),
    ("02", "The right people in the room:", "those who run the work day to day, and a decision maker for the checkpoint."),
    ("03", "The evidence you already have:", "procedures, process maps, slides, notes or BPMN files, and a system event log if one is available."),
    ("04", "Rough figures:", "volumes, effort, waiting times and costs. Estimates are fine; we label them as assumptions and test which ones matter."),
], row_h=0.72, gap=0.3)
picture(s, A / "illustrations" / "team-desk.png", 8.45, 2.0, w=4.22)
notes(s, NOTES[13])

# ---------------- 15 close (template section slide) ----------------
for sh in S_CLOSE.shapes:
    if sh.is_placeholder and sh.placeholder_format.idx == 0:
        set_text(sh, "Which process should we put under the AIxRay?")
        for r in sh.text_frame.paragraphs[0].runs:
            r.font.size = Pt(32)
    elif sh.is_placeholder and sh.placeholder_format.idx == 1:
        set_text(sh, "Choose a candidate process, confirm participants and dates, and share what you have so we can prepare a first model.")
notes(S_CLOSE, NOTES[14])
notes(S_ABOUT, "Before we finish, a word about us. Sogeti is part of the Capgemini Group, and we focus on making technology simple and impactful, with AI-driven solutions across quality engineering and testing, data and cloud. AIxRay is one of the ways we help clients turn ideas into action quickly.")
notes(S_LOGO, "Thank you for your time. I look forward to choosing the first process with you.")

# ---------------- order: title, new slides, close, about, logo ----------------
ids = list(sld_ids)
by_id = {int(e.get("id")): e for e in ids}
first = [by_id[S_TITLE.slide_id]]
tail = [by_id[S_CLOSE.slide_id], by_id[S_ABOUT.slide_id], by_id[S_LOGO.slide_id]]
middle = [e for e in ids if e not in first and e not in tail]
for e in ids:
    sld_ids.remove(e)
for e in first + middle + tail:
    sld_ids.append(e)

# renumber slide and notes parts in final order so kept template slides and new slides never share a partname
from pptx.opc.packuri import PackURI
for i, sl in enumerate(prs.slides, 1):
    sl.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")
    if sl.has_notes_slide:
        sl.notes_slide.part.partname = PackURI(f"/ppt/notesSlides/notesSlide{i}.xml")

prs.core_properties.title = "AIxRay in your Jump Start workshop"
prs.core_properties.author = "Sogeti"
prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
